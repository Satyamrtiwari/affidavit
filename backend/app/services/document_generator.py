"""
Document Generator Service — Stage 4 of the pipeline.

Takes CaseEntities (structured JSON) and produces:
1. Plain text version (via Jinja2 template) — used for evaluation
2. Formatted .docx version (via python-docx) — the deliverable

The Jinja2 template handles content and structure.
python-docx handles formatting (bold, caps, centered, justified).
"""

import logging
from pathlib import Path
from io import BytesIO

from jinja2 import Environment, FileSystemLoader
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.config import TEMPLATES_DIR, OUTPUTS_DIR
from app.models.entities import CaseEntities

logger = logging.getLogger(__name__)


def _get_jinja_env() -> Environment:
    """Create Jinja2 environment with the templates directory."""
    return Environment(
        loader=FileSystemLoader(str(TEMPLATES_DIR)),
        trim_blocks=True,
        lstrip_blocks=True,
        keep_trailing_newline=False,
    )


def generate_text(entities: CaseEntities) -> str:
    """
    Generate the affidavit as plain text using Jinja2 template.

    This plain text version is used for:
    - Evaluation/validation checks
    - Preview display
    - Debugging

    Args:
        entities: Validated CaseEntities from extraction stage

    Returns:
        Plain text of the generated affidavit
    """
    env = _get_jinja_env()
    template = env.get_template("affidavit_reply.j2")

    # Build template context from entities
    context = entities.model_dump()

    # Add computed properties
    context["clean_jurisdiction"] = entities.clean_jurisdiction
    context["clean_prayer_points"] = entities.clean_prayer_points
    context["paragraph_count"] = entities.paragraph_count
    context["jurat_verb"] = entities.jurat_verb
    context["ordinal_date"] = entities.ordinal_date

    # Add clean entity properties
    context["petitioner"]["clean_through"] = entities.petitioner.clean_through
    context["petitioner"]["clean_address"] = entities.petitioner.clean_address
    for r_dict, r_obj in zip(context["respondents"], entities.respondents):
        r_dict["clean_through"] = r_obj.clean_through
        r_dict["clean_address"] = r_obj.clean_address
    context["deponent"]["clean_address"] = entities.deponent.clean_address

    # Add formatted reply points
    context["formatted_points"] = [
        {"point_number": p.point_number, "text": p.format_for_display(entities)}
        for p in entities.reply_points
    ]

    rendered = template.render(**context)

    # Clean up excessive blank lines from template rendering
    import re
    rendered = re.sub(r'\n{3,}', '\n\n', rendered)
    rendered = rendered.strip()

    logger.info(f"Generated plain text: {len(rendered)} chars, {entities.paragraph_count} paragraphs")
    return rendered


def generate_docx(entities: CaseEntities, output_path: Path = None) -> BytesIO:
    """
    Generate a formatted .docx file from CaseEntities.

    Applies proper legal document formatting:
    - Bold, ALL CAPS, centered for headings
    - Justified body text
    - Bold paragraph numbers
    - Right-aligned DEPONENT
    - Proper spacing and fonts

    Args:
        entities: Validated CaseEntities from extraction stage
        output_path: Optional path to save the .docx file

    Returns:
        BytesIO buffer containing the .docx file
    """
    doc = Document()

    # ─── Default style setup ──────────────────────────────────────────────
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15
    try:
        from docx.oxml.ns import qn
        rFonts = style.element.rPr.get_or_add_rFonts()
        rFonts.set(qn("w:ascii"), "Times New Roman")
        rFonts.set(qn("w:hAnsi"), "Times New Roman")
        rFonts.set(qn("w:cs"), "Times New Roman")
    except Exception:
        pass

    # ─── Helper functions ─────────────────────────────────────────────────
    def add_centered_bold_caps(text: str, space_before: int = 0, space_after: int = 6):
        """Add a centered, bold, ALL CAPS line."""
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(12)
        return p

    def add_normal_text(text: str, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY):
        """Add a normal justified paragraph."""
        p = doc.add_paragraph()
        p.alignment = alignment
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.size = Pt(12)
        return p

    def add_numbered_paragraph(number: int, text: str):
        """Add a numbered paragraph with bold number."""
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(8)
        # Bold number
        num_run = p.add_run(f"{number}. ")
        num_run.bold = True
        num_run.font.size = Pt(12)
        # Normal text
        text_run = p.add_run(text)
        text_run.font.size = Pt(12)
        return p

    def add_right_aligned(text: str, bold: bool = False, caps: bool = False):
        """Add a right-aligned line."""
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(4)
        display_text = text.upper() if caps else text
        run = p.add_run(display_text)
        run.bold = bold
        run.font.size = Pt(12)
        return p

    # ─── Part 1: Forum Heading ────────────────────────────────────────────
    add_centered_bold_caps(
        f"IN THE HIGH COURT OF JUDICATURE AT {entities.forum_city}",
        space_before=12
    )

    # ─── Part 2: Jurisdiction ─────────────────────────────────────────────
    add_centered_bold_caps(entities.clean_jurisdiction)

    # ─── Part 3: Case Number ─────────────────────────────────────────────
    add_centered_bold_caps(
        f"{entities.case_type} NO. {entities.case_number} OF {entities.year}",
        space_after=12
    )

    # ─── Part 4: Cause Title ──────────────────────────────────────────────
    # Petitioner
    pet = entities.petitioner
    pet_text = pet.name
    if pet.description:
        pet_text += f", {pet.description.rstrip('.')}"
    if pet.clean_address and (not pet.description or pet.clean_address not in pet.description):
        pet_text += f",\nhaving registered office at {pet.clean_address}"
    if pet.clean_through:
        pet_text += f"\n{pet.clean_through}."
    else:
        pet_text = pet_text.rstrip('.') + "."

    add_normal_text(pet_text, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    add_right_aligned("...Petitioner")

    # VERSUS
    add_centered_bold_caps("VERSUS", space_before=6, space_after=6)

    # Respondents
    for resp in entities.respondents:
        resp_text = f"{resp.respondent_number}. {resp.name.rstrip('.')}"
        if resp.description:
            resp_text += f", {resp.description.rstrip('.')}"
        if resp.clean_through:
            resp_text += f",\n   {resp.clean_through}"
        if resp.clean_address and (not resp.description or resp.clean_address not in resp.description):
            addr_lead = "having office at" if resp.is_organisation else "residing at"
            resp_text += f",\n   {addr_lead} {resp.clean_address}"
        resp_text = resp_text.rstrip('.') + "."

        add_normal_text(resp_text, alignment=WD_ALIGN_PARAGRAPH.LEFT)
        add_right_aligned(f"...Respondent No.{resp.respondent_number}")

    # ─── Part 5: Affidavit Title ──────────────────────────────────────────
    add_centered_bold_caps(
        f"AFFIDAVIT IN REPLY ON BEHALF OF RESPONDENT NO. {entities.filing_respondent_number}",
        space_before=12,
        space_after=12
    )

    # ─── Part 6: Deponent Clause ──────────────────────────────────────────
    dep = entities.deponent
    if dep.is_organisation_representative:
        deponent_text = (
            f"I, {dep.name}, having office at {dep.clean_address}, "
            f"the {dep.designation} of the Respondent No.{entities.filing_respondent_number} "
            f"above named, do hereby {dep.verification_verb} and state as under:"
        )
    else:
        parts = [f"I, {dep.name}"]
        if dep.age:
            parts.append(f"Age {dep.age.rstrip('.')}")
        if dep.occupation:
            parts.append(f"Occupation: {dep.occupation.rstrip('.')}")
        parts_text = ", ".join(parts)
        deponent_text = (
            f"{parts_text}, residing at {dep.clean_address}, "
            f"the Respondent No.{entities.filing_respondent_number} above named, "
            f"do hereby {dep.verification_verb} and state as under:"
        )

    add_normal_text(deponent_text)

    # ─── Part 7: Numbered Paragraphs ──────────────────────────────────────
    for point in entities.reply_points:
        para_text = point.format_for_display(entities)
        add_numbered_paragraph(point.point_number, para_text)

    # ─── Part 8: Prayer ───────────────────────────────────────────────────
    add_centered_bold_caps("PRAYER", space_before=12)

    prayer_intro = "I therefore respectfully pray that this Hon'ble Court may be pleased to:"
    add_normal_text(prayer_intro)

    letters = "abcdefghij"
    prayers = entities.clean_prayer_points
    for i, prayer in enumerate(prayers):
        letter = letters[i] if i < len(letters) else str(i + 1)
        suffix = "; and" if i < len(prayers) - 1 else "."
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.space_after = Pt(4)
        # Bold letter
        letter_run = p.add_run(f"({letter}) ")
        letter_run.bold = True
        letter_run.font.size = Pt(12)
        # Prayer text
        clean_p = prayer.strip().rstrip(".;")
        text_run = p.add_run(f"{clean_p}{suffix}")
        text_run.font.size = Pt(12)

    # ─── Part 9: Jurat ───────────────────────────────────────────────────
    doc.add_paragraph()  # Spacer
    add_normal_text(
        f"{entities.jurat_verb} at {entities.attestation_place}",
        alignment=WD_ALIGN_PARAGRAPH.LEFT
    )
    add_normal_text(
        f"On this {entities.ordinal_date}",
        alignment=WD_ALIGN_PARAGRAPH.LEFT
    )

    doc.add_paragraph()  # Spacer
    add_right_aligned("DEPONENT", bold=True, caps=True)
    add_normal_text("Before Me", alignment=WD_ALIGN_PARAGRAPH.LEFT)

    # ─── Part 10: Verification ────────────────────────────────────────────
    add_centered_bold_caps("VERIFICATION", space_before=12)

    verification_text = (
        f"I, {dep.name}, the Deponent above named, do hereby verify that the "
        f"contents of paragraphs 1 to {entities.paragraph_count} and the Prayer above "
        f"are true and correct to my knowledge and belief and that nothing material "
        f"has been concealed therefrom."
    )
    add_normal_text(verification_text)

    add_normal_text(
        f"Verified at {entities.attestation_place} on this {entities.ordinal_date}."
    )

    doc.add_paragraph()  # Spacer
    add_right_aligned("DEPONENT", bold=True, caps=True)

    # ─── Advocate Block (if present) ──────────────────────────────────────
    if entities.advocate_firm:
        doc.add_paragraph()  # Spacer
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        firm_run = p.add_run(entities.advocate_firm.upper())
        firm_run.bold = True
        firm_run.font.size = Pt(12)

        add_normal_text(
            f"Advocates for the Respondent No.{entities.filing_respondent_number}.",
            alignment=WD_ALIGN_PARAGRAPH.LEFT
        )

    # ─── Save ─────────────────────────────────────────────────────────────
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    if output_path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        logger.info(f"Saved .docx to: {output_path}")

    logger.info("Generated formatted .docx successfully")
    return buffer


def generate_all(
    entities: CaseEntities,
    save_docx: bool = True,
    output_filename: str = "generated_affidavit.docx"
) -> tuple[str, BytesIO]:
    """
    Generate both plain text and .docx versions.

    Args:
        entities: Validated CaseEntities
        save_docx: Whether to save the .docx to the outputs directory
        output_filename: Filename for the generated .docx file

    Returns:
        Tuple of (plain_text, docx_buffer)
    """
    plain_text = generate_text(entities)

    output_path = OUTPUTS_DIR / output_filename if save_docx else None
    docx_buffer = generate_docx(entities, output_path=output_path)

    return plain_text, docx_buffer

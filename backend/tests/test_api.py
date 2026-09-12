"""
Integration tests for FastAPI endpoints.
"""

import pytest
from fastapi.testclient import TestClient
from app.main import app
from app.config import REFERENCE_DOCS_DIR


@pytest.fixture
def client():
    return TestClient(app)


def test_health_check(client):
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json()["status"] == "healthy"


def test_generate_endpoint_requires_file(client):
    response = client.post("/api/generate")
    # Missing required form-data file
    assert response.status_code == 422


def test_generate_endpoint_with_sample_pdf(client):
    case_pdf = REFERENCE_DOCS_DIR / "03_Case_Information.pdf"
    if not case_pdf.exists():
        pytest.skip("Reference PDF not found")

    with open(case_pdf, "rb") as f:
        response = client.post(
            "/api/generate",
            files={"case_info": ("03_Case_Information.pdf", f, "application/pdf")},
        )

    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert "entities" in data
    assert "evaluation" in data
    assert data["evaluation"]["overall_score"] >= 80.0
    assert "generated_03_Case_Information.docx" in data["files"]["affidavit"]

    # Verify linear multi-agent execution trajectory
    assert "agent_trajectory" in data
    trajectory = data["agent_trajectory"]
    assert len(trajectory) == 5
    expected_agents = [
        "ExtractorAgent",
        "PreGenerationGuardAgent",
        "DrafterAgent",
        "EvaluatorAgent",
        "ExporterAgent",
    ]
    assert [step["agent_name"] for step in trajectory] == expected_agents
    for step in trajectory:
        assert step["status"] == "success"
        assert step["duration_ms"] >= 0


def test_download_endpoints(client):
    # Test auto-named docx download without specifying filename
    response_docx = client.get("/api/download/affidavit")
    assert response_docx.status_code in (200, 404)
    if response_docx.status_code == 200:
        cd = response_docx.headers.get("content-disposition", "")
        assert "generated_" in cd and cd.endswith('.docx"')

    # Test auto-named json report download without specifying filename
    response_report = client.get("/api/download/report?format=json")
    assert response_report.status_code in (200, 404)
    if response_report.status_code == 200:
        cd = response_report.headers.get("content-disposition", "")
        assert "evaluation_report_" in cd and cd.endswith('.json"')

    # Test auto-named md report download
    response_report_md = client.get("/api/download/report?format=md")
    assert response_report_md.status_code in (200, 404)
    if response_report_md.status_code == 200:
        cd = response_report_md.headers.get("content-disposition", "")
        assert "evaluation_report_" in cd and cd.endswith('.md"')

    # Test resolving by stem name e.g. "03_Case_Information"
    response_stem = client.get("/api/download/affidavit?filename=03_Case_Information")
    assert response_stem.status_code in (200, 404)
    if response_stem.status_code == 200:
        assert 'filename="generated_03_Case_Information.docx"' in response_stem.headers.get("content-disposition", "")


def test_generate_endpoint_with_empty_reference_doc_field(client):
    case_pdf = REFERENCE_DOCS_DIR / "03_Case_Information.pdf"
    if not case_pdf.exists():
        pytest.skip("Reference PDF not found")

    with open(case_pdf, "rb") as f:
        # Simulate Swagger sending empty string for optional reference_doc
        response = client.post(
            "/api/generate",
            files={"case_info": ("03_Case_Information.pdf", f, "application/pdf")},
            data={"reference_doc": ""}
        )

    assert response.status_code == 200
    assert response.json()["success"] is True


def test_extract_entities_with_docx(client):
    import docx
    from io import BytesIO

    doc = docx.Document()
    doc.add_paragraph("IN THE HIGH COURT OF JUDICATURE AT BOMBAY")
    doc.add_paragraph("WRIT PETITION NO. 4521 OF 2024")
    doc.add_paragraph("PETITIONER: Rajesh Sharma")
    doc.add_paragraph("RESPONDENT: State of Maharashtra")
    doc.add_paragraph("ADVOCATE FOR RESPONDENT: Adv. S. Patil")
    doc.add_paragraph("DATE OF NOTICE: 15-08-2024")
    
    buf = BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    response = client.post(
        "/api/extract-entities",
        files={"case_info": ("test_case.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert "Rajesh" in data["entities"]["petitioner"]["name"]


def test_generate_endpoint_with_docx(client):
    import docx
    from io import BytesIO
    from app.services.pdf_parser import parse_pdf

    # Read 03_Case_Information.pdf to build realistic docx
    pdf_path = REFERENCE_DOCS_DIR / "03_Case_Information.pdf"
    if pdf_path.exists():
        content = parse_pdf(pdf_path)
    else:
        content = "COURT: HIGH COURT OF JUDICATURE AT BOMBAY\nCASE: WRIT PETITION NO. 1847 OF 2026\nPETITIONER: Arvind Rajan\nRESPONDENT: Sunrise Housing Private Limited\nDEPONENT: Director"

    doc = docx.Document()
    for line in content.split("\n"):
        if line.strip():
            doc.add_paragraph(line.strip())

    buf = BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    response = client.post(
        "/api/generate",
        files={"case_info": ("case_info_test.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert "entities" in data
    assert "evaluation" in data
    assert data["evaluation"]["overall_score"] >= 80.0
    assert "generated_case_info_test.docx" in data["files"]["affidavit"]
    assert len(data["agent_trajectory"]) == 5


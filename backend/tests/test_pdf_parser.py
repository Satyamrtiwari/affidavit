"""
Unit tests for PDF Parser service.
"""

import pytest
from pathlib import Path
from io import BytesIO
import docx

from app.services.pdf_parser import clean_text, parse_pdf, parse_document, extract_text_from_docx, extract_text_from_txt
from app.config import REFERENCE_DOCS_DIR


def test_clean_text_removes_broken_words():
    raw = "IN THE HIGH COUR T OF JUDICA TURE AT BOMBA Y"
    cleaned = clean_text(raw)
    assert "COURT" in cleaned
    assert "JUDICATURE" in cleaned
    assert "BOMBAY" in cleaned


def test_clean_text_normalizes_whitespace():
    raw = "Line 1    \n\n\n\n\nLine 2     with   spaces"
    cleaned = clean_text(raw)
    assert "Line 1\n\nLine 2 with spaces" == cleaned


def test_parse_pdf_reference_files():
    case_pdf = REFERENCE_DOCS_DIR / "03_Case_Information.pdf"
    if case_pdf.exists():
        text = parse_pdf(case_pdf)
        assert len(text) > 500
        assert "Sunrise Housing Private Limited" in text
        assert "Arvind Rajan" in text


def test_parse_docx_in_memory():
    # Build in-memory docx document
    doc = docx.Document()
    doc.add_heading("IN THE HIGH COURT OF JUDICATURE AT BOMBAY", level=1)
    doc.add_paragraph("PETITIONER: Rajesh Sharma")
    doc.add_paragraph("RESPONDENT: State of Maharashtra")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Case Number"
    table.cell(0, 1).text = "WP 4521 of 2024"
    table.cell(1, 0).text = "Date of Filing"
    table.cell(1, 1).text = "15-08-2024"

    buf = BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    # Test parse_document with filename
    text = parse_document(docx_bytes, filename="case_info.docx")
    assert "IN THE HIGH COURT OF JUDICATURE AT BOMBAY" in text
    assert "Rajesh Sharma" in text
    assert "WP 4521 of 2024" in text

    # Test parse_pdf fallback / auto-sniffing
    text_via_pdf_alias = parse_pdf(docx_bytes, filename="case_info.docx")
    assert "Rajesh Sharma" in text_via_pdf_alias


def test_parse_txt_file():
    txt_content = "PETITIONER: John Doe\nRESPONDENT: Jane Doe\nCASE TYPE: Writ Petition"
    txt_bytes = txt_content.encode("utf-8")
    
    text = parse_document(txt_bytes, filename="case_info.txt")
    assert "John Doe" in text
    assert "Writ Petition" in text

